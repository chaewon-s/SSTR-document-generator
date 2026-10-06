from __future__ import annotations
import re
from docx import Document

# Remove a trailing Jira issue tag like "[KKZ-8606]" or "[ KKZ-8393 ]"
_ISSUE_TAG = re.compile(r"\s*\[\s*[A-Z]{2,}-\d+\s*\]\s*$")
# A valid HUD SW build string starts with a 4-char release info block (e.g. "2691.")
# or the legacy "HUD_C." prefix.
_SW_PATTERN = re.compile(r"(HUD_C\.|[0-9A-Za-z]{4}\.[A-Z0-9]{2,4}\.(CS|CW)\.HUD)")


def _clean_sw(sw: str) -> str:
    """Strip the trailing [KKZ-xxxx] tag and surrounding whitespace."""
    return _ISSUE_TAG.sub("", sw).strip()


def _iter_table_rows(doc):
    for tbl in doc.tables:
        for row in tbl.rows:
            yield [c.text.strip() for c in row.cells]


def load_sw_versions(path: str) -> list:
    """
    Read BOTH SW version tables (legacy variants + 3rd-gen variants).
    Vehicle cell is merged across phases, so the last non-empty value is inherited.
    Rows whose SW string does not match the build-string pattern are skipped
    (this drops header rows and struck-through obsolete entries that lost markup).
    """
    doc = Document(path)
    out, last = [], ""
    for cells in _iter_table_rows(doc):
        if len(cells) != 3:
            continue
        vehicle, phase, sw = cells
        if not _SW_PATTERN.search(sw):
            continue
        if vehicle:
            last = vehicle
        out.append({
            "vehicle": last,
            "phase": phase,
            "sw_version": _clean_sw(sw),
        })
    return out


def find_sw_version(path: str, vehicle: str, phase: str):
    """Look up one build string by vehicle + phase (case/space-insensitive)."""
    n = lambda s: s.upper().replace(" ", "")
    for r in load_sw_versions(path):
        if n(r["vehicle"]) == n(vehicle) and n(r["phase"]) == n(phase):
            return r["sw_version"]
    return None


def list_phases(path: str, vehicle: str) -> list:
    """All phases available for a given vehicle (for a UI dropdown)."""
    n = lambda s: s.upper().replace(" ", "")
    return [r["phase"] for r in load_sw_versions(path) if n(r["vehicle"]) == n(vehicle)]


def load_latest_revision(path: str):
    """
    Return the most recent REVISION HISTORY row that carries a version number.
    NOTE: this is the LATEST revision; the tester may need an earlier one
    (e.g. SSTR shows 2.08 while latest is 2.09), so the UI must allow editing.
    """
    doc = Document(path)
    latest = None
    for cells in _iter_table_rows(doc):
        if len(cells) < 4:
            continue
        ver = cells[0].strip()
        if ver and ver[0].isdigit() and "." in ver:
            latest = {"ver": ver, "contents": cells[1], "released": cells[3]}
    return latest
