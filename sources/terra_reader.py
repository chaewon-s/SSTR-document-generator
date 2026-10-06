from __future__ import annotations
from docx import Document


def _cell_pairs(tbl):
    pairs = []
    for row in tbl.rows:
        cells = [c.text.strip() for c in row.cells]
        for i, c in enumerate(cells[:-1]):
            if c.endswith(":") and cells[i + 1]:
                pairs.append((c.rstrip(":").strip(), cells[i + 1]))
    return pairs


def load_terra_fields(path: str) -> dict:
    doc = Document(path)
    out = {"metrics": {}, "test_matrix": [], "sw_version": None, "test_period": None}
    for tbl in doc.tables:
        for label, value in _cell_pairs(tbl):
            L = label.lower()
            if L.startswith("sw-version"):
                out["sw_version"] = value.splitlines()[0].strip()
            elif "test execution time" in L:
                out["metrics"]["Test execution time"] = value
            elif "no. of planned test cases" in L:
                out["metrics"]["No. of Planned Test Cases"] = value
            elif "no. of executed test cases" in L:
                out["metrics"]["No. of Executed Test Cases"] = value
            elif "percentage of test execution" in L:
                out["metrics"]["Percentage of Test Execution"] = value
    for p in doc.paragraphs:
        if p.text.strip().lower().startswith("test period"):
            out["test_period"] = p.text.split(":", 1)[-1].strip()
    for tbl in doc.tables:
        header = [c.text.strip() for c in tbl.rows[0].cells]
        if any("Test Objective" in h for h in header):
            out["test_matrix"] = [
                [c.text.strip() for c in r.cells]
                for r in tbl.rows
                if r.cells[0].text.strip()
            ]
            break
    return out
