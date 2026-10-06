from __future__ import annotations
from docx import Document


def _set_value_after_label(doc, label, value):
    target = label.lower().rstrip(":").strip()
    for tbl in doc.tables:
        for row in tbl.rows:
            cells = row.cells
            for i, c in enumerate(cells[:-1]):
                if c.text.strip().lower().rstrip(":").strip() == target:
                    if not cells[i + 1].text.strip():
                        cells[i + 1].text = value
                    return True
    return False


def _apply_metrics(doc, metrics: dict):
    """teRA Metrics 값을 Metrics 표에 복사 (B)."""
    for key, val in metrics.items():
        _set_value_after_label(doc, key, val)


def _ensure_test_matrix(doc, matrix: list):
    """
    teRA D.Test Matrix 를 템플릿의 동일 표에 복사 (B).
    헤더에 'Test Objective' 가 있는 표를 찾아, 데이터 행이 비어 있으면 추가.
    """
    if not matrix:
        return
    for tbl in doc.tables:
        header = [c.text.strip() for c in tbl.rows[0].cells]
        if not any("Test Objective" in h for h in header):
            continue
        existing = sum(1 for r in tbl.rows[1:] if r.cells[0].text.strip())
        if existing > 0:
            return  # 이미 채워져 있으면 건드리지 않음
        for data_row in matrix[1:]:  # matrix[0] 은 헤더
            cells = tbl.add_row().cells
            for i, val in enumerate(data_row):
                if i < len(cells):
                    cells[i].text = val
        return


def _set_test_period(doc, period: str):
    if not period:
        return
    for p in doc.paragraphs:
        txt = p.text.strip().lower()
        if txt.startswith("test period") and ":" not in p.text.split("period")[-1]:
            p.add_run(f" {period}")
            return


def fill_sstr(template_path: str, out_path: str, data: dict,
              terra: dict | None = None) -> str:
    """
    data  : 라벨->값 (Customer, Part Number, Test Type 등)
    terra : terra_reader.load_terra_fields() 결과 (B: metrics/matrix/period 복사)
    """
    doc = Document(template_path)
    missed = []
    for label, value in data.items():
        if value in (None, ""):
            continue
        if not _set_value_after_label(doc, label, str(value)):
            missed.append(label)

    if terra:
        _apply_metrics(doc, terra.get("metrics", {}))
        _ensure_test_matrix(doc, terra.get("test_matrix", []))
        _set_test_period(doc, terra.get("test_period"))

    doc.save(out_path)
    if missed:
        print("[경고] 채우지 못한 라벨:", missed)
    return out_path
